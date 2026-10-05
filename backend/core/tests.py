import io
import sqlite3
import tempfile
from contextlib import closing
from datetime import timedelta
from pathlib import Path
from unittest import mock

from django.conf import settings as django_settings
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.management import call_command
from django.test import TransactionTestCase, override_settings
from rest_framework.test import APITestCase

from core import services as svc
from core.models import Attendance, AuditLog, DailyPayment, ExpenseCategory, Opinion, Settings, Worker


class RulesTest(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user('m', password='x')
        self.client.force_authenticate(self.user)
        Settings.load()
        ExpenseCategory.objects.create(key='gesi', name_sw='Gesi', name_en='Gas')
        self.today = svc.today()
        long_ago = self.today - timedelta(days=60)
        mk = lambda name, role, pay, rate, off: Worker.objects.create(
            name=name, role=role, pay_type=pay, rate=rate, day_off=off, joined_on=long_ago)
        # day off is set to a weekday that is never "today" so marking is required
        off = (svc.js_dow(self.today) + 3) % 7
        self.cashier = mk('Rehema', 'keshia', 'monthly', 300000, off)
        self.cook = mk('Juma', 'mpishi', 'daily', 10000, off)
        self.d = self.today.isoformat()

    def sales(self, banda_cash=100000, cashier=None):
        r = self.client.put(f'/api/days/{self.d}/sales/', {
            'banda': {'cash': banda_cash, 'mobile': 50000, 'cashier_id': (cashier or self.cashier).id},
            'mgahawa': {'cash': 200000, 'mobile': 0, 'cashier_id': self.cashier.id}}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        return r.data

    def mark_all(self):
        r = self.client.put(f'/api/days/{self.d}/attendance/', {str(self.cashier.id): 'present', str(self.cook.id): 'late'}, format='json')
        self.assertEqual(r.status_code, 200)

    def test_expected_cash_and_difference(self):
        self.sales()
        self.client.post('/api/expenses/', {'category': 'gesi', 'amount': 25000, 'reason': 'gas', 'paid_from': 'droo', 'section': 'banda'}, format='json')
        r = self.client.put(f'/api/days/{self.d}/cash-count/', {'banda': 50000 + 75000 - 2000, 'mgahawa': 80000 + 200000}, format='json')
        b = r.data['sections']['banda']
        self.assertEqual(b['expected'], 100000 - 25000)
        self.assertEqual(b['difference'], 73000 - 75000)  # counted 123000 - float 50000 - expected 75000 = -2000
        self.assertEqual(r.data['sections']['mgahawa']['difference'], 0)

    def test_close_writes_shortage_and_locks(self):
        self.sales()
        self.client.put(f'/api/days/{self.d}/cash-count/', {'banda': 50000 + 98000, 'mgahawa': 280000}, format='json')
        self.assertEqual(self.client.post(f'/api/days/{self.d}/close/').status_code, 409)  # attendance missing
        self.mark_all()
        r = self.client.post(f'/api/days/{self.d}/close/')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual([(s['worker_id'], s['amount']) for s in r.data['shortages']], [(self.cashier.id, 2000)])
        self.assertEqual(self.client.put(f'/api/days/{self.d}/sales/', {'banda': {'cash': 1}}, format='json').status_code, 409)
        acct = self.client.get(f'/api/workers/{self.cashier.id}/account/').data
        self.assertEqual(acct['shortages'], 2000)
        self.assertEqual(acct['net'], acct['base'] - 2000)

    def test_waive_and_salary_lock(self):
        self.sales()
        self.client.put(f'/api/days/{self.d}/cash-count/', {'banda': 50000 + 98000, 'mgahawa': 280000}, format='json')
        self.mark_all()
        sh = self.client.post(f'/api/days/{self.d}/close/').data['shortages'][0]
        ym = svc.ym_of(self.today)
        self.assertEqual(self.client.patch(f'/api/shortages/{sh["id"]}/', {'status': 'waived'}, format='json').status_code, 200)
        self.assertEqual(self.client.get(f'/api/workers/{self.cashier.id}/account/').data['shortages'], 0)
        self.assertEqual(self.client.post(f'/api/salary/{ym}/approve/').status_code, 200)
        self.assertEqual(self.client.patch(f'/api/shortages/{sh["id"]}/', {'status': 'applied'}, format='json').status_code, 409)
        self.assertEqual(self.client.post('/api/advances/', {'worker_id': self.cook.id, 'amount': 1000}, format='json').status_code, 409)

    def test_change_password(self):
        r = self.client.post('/api/auth/change-password/', {'current_password': 'wrong', 'new_password': 'NewPass123!'}, format='json')
        self.assertEqual((r.status_code, r.data['code']), (409, 'wrong_password'))
        r = self.client.post('/api/auth/change-password/', {'current_password': 'x', 'new_password': '123'}, format='json')
        self.assertEqual(r.status_code, 400)
        r = self.client.post('/api/auth/change-password/', {'current_password': 'x', 'new_password': 'NewPass123!'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password('NewPass123!'))

    def test_demo_mode_blocks_password_change(self):
        with self.settings(DEMO_MODE=True):
            r = self.client.post('/api/auth/change-password/', {'current_password': 'x', 'new_password': 'NewPass123!'}, format='json')
        self.assertEqual((r.status_code, r.data['code']), (409, 'demo_mode'))
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password('x'))

    def test_category_delete(self):
        r = self.client.post('/api/categories/', {'name_sw': 'Usafi', 'name_en': 'Cleaning'}, format='json')
        cat_id = r.data['id']
        self.assertEqual(self.client.delete(f'/api/categories/{cat_id}/').status_code, 204)
        self.client.post('/api/categories/', {'name_sw': 'Usafi', 'name_en': 'Cleaning'}, format='json')
        cat_id2 = ExpenseCategory.objects.get(name_en='Cleaning').id
        self.client.post('/api/expenses/', {'category': ExpenseCategory.objects.get(id=cat_id2).key, 'amount': 5000, 'reason': 'Soap', 'paid_from': 'simu'}, format='json')
        r = self.client.delete(f'/api/categories/{cat_id2}/')
        self.assertEqual((r.status_code, r.data['code']), (409, 'category_in_use'))
        self.assertTrue(ExpenseCategory.objects.filter(id=cat_id2).exists())

    def test_no_day_off_worker(self):
        w = Worker.objects.create(name='Siku Zote', role='mhudumu', pay_type='daily', rate=5000, day_off=None,
                                   joined_on=self.today - timedelta(days=5))
        for i in range(5):
            key, locked = svc.status_on(w, self.today - timedelta(days=i))
            self.assertFalse(locked)  # never auto-filled as a day off

        r = self.client.post('/api/workers/', {'name': 'Kila Siku', 'role': 'mpishi', 'pay_type': 'daily', 'rate': 7000, 'day_off': None}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertIsNone(r.data['day_off'])
        wid = r.data['id']

        r = self.client.patch(f'/api/workers/{wid}/', {'day_off': 2}, format='json')
        self.assertEqual(r.data['day_off'], 2)
        r = self.client.patch(f'/api/workers/{wid}/', {'day_off': None}, format='json')
        self.assertIsNone(r.data['day_off'])

    def test_joined_on_registration(self):
        future = self.today + timedelta(days=3)
        r = self.client.post('/api/workers/', {'name': 'Mpya Baadaye', 'role': 'mhudumu', 'pay_type': 'daily', 'rate': 6000, 'joined_on': future.isoformat()}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data['joined_on'], future.isoformat())
        wid = r.data['id']
        today_payload = self.client.get(f'/api/days/{self.today}/').data
        self.assertNotIn(wid, [a['worker_id'] for a in today_payload['attendance']])  # not active yet, hasn't started

        past = self.today - timedelta(days=10)
        r = self.client.patch(f'/api/workers/{wid}/', {'joined_on': past.isoformat()}, format='json')
        self.assertEqual(r.data['joined_on'], past.isoformat())
        today_payload = self.client.get(f'/api/days/{self.today}/').data
        self.assertIn(wid, [a['worker_id'] for a in today_payload['attendance']])  # now active

        r = self.client.post('/api/workers/', {'name': 'Bila Tarehe', 'role': 'mhudumu', 'pay_type': 'daily', 'rate': 6000}, format='json')
        self.assertEqual(r.data['joined_on'], self.today.isoformat())  # defaults to today

    def test_monthly_salary_prorated_for_partial_month(self):
        # Use the previous, already-finished month so "today" truncation can't interfere.
        first_this, _ = svc.month_bounds(svc.ym_of(self.today))
        prev_last = first_this - timedelta(days=1)
        ym = svc.ym_of(prev_last)
        first, last = svc.month_bounds(ym)
        days_in_month = (last - first).days + 1
        join_day = first + timedelta(days=days_in_month - 5)  # registered with 5 days left in the month
        w = Worker.objects.create(name='Mwezi Nusu', role='mhudumu', pay_type='monthly', rate=300000,
                                   day_off=None, joined_on=join_day)
        d = join_day
        while d <= last:
            Attendance.objects.create(worker=w, date=d, status='present')
            d += timedelta(days=1)
        rules = Settings.load().rules
        f = svc.salary_figures(w, ym, rules)
        employed_days = (last - join_day).days + 1
        expected = max(0, 300000 - round(300000 / 30 * (days_in_month - employed_days)))
        self.assertEqual(f['base'], expected)
        self.assertLess(f['base'], 300000)  # the registration date actually prorates the salary

    def test_monthly_salary_prorated_for_mid_month_removal(self):
        first_this, _ = svc.month_bounds(svc.ym_of(self.today))
        prev_last = first_this - timedelta(days=1)
        ym = svc.ym_of(prev_last)
        first, last = svc.month_bounds(ym)
        days_in_month = (last - first).days + 1
        removed = first + timedelta(days=4)  # only worked the first 5 days of the month
        w = Worker.objects.create(name='Aliondoka Mapema', role='mhudumu', pay_type='monthly', rate=300000,
                                   day_off=None, joined_on=first - timedelta(days=200), active=False, removed_on=removed)
        d = first
        while d <= removed:
            Attendance.objects.create(worker=w, date=d, status='present')
            d += timedelta(days=1)
        rules = Settings.load().rules
        f = svc.salary_figures(w, ym, rules)
        employed_days = (removed - first).days + 1
        expected = max(0, 300000 - round(300000 / 30 * (days_in_month - employed_days)))
        self.assertEqual(f['base'], expected)
        self.assertLess(f['base'], 300000)

    def test_monthly_salary_not_prorated_just_because_month_is_not_over(self):
        ym = svc.ym_of(self.today)
        first, last = svc.month_bounds(ym)
        # self.cashier joined long before this month and has not been removed.
        self.assertEqual(svc.month_employed_days(self.cashier, first, last), (last - first).days + 1)

    def test_unpaid_salary_carries_forward(self):
        # the monthly-paid cashier: daily-paid workers are paid each day and are not on the monthly list
        self.sales()
        self.client.put(f'/api/days/{self.d}/cash-count/', {'banda': 50000 + 98000, 'mgahawa': 280000}, format='json')
        self.mark_all()
        self.client.post(f'/api/days/{self.d}/close/')
        ym = svc.ym_of(self.today)
        r = self.client.post(f'/api/salary/{ym}/approve/')
        self.assertEqual(r.status_code, 200)
        net = next(l['net'] for l in r.data['lines'] if l['worker']['id'] == self.cashier.id)
        self.assertGreater(net, 0)

        # approved but not marked paid -> it's money the company owes the cashier, visible right away
        w = self.client.get(f'/api/workers/{self.cashier.id}/').data
        self.assertEqual(w['company_owes'], net)
        today_payload = self.client.get(f'/api/days/{self.d}/').data
        self.assertIn(self.cashier.id, [u['worker_id'] for u in today_payload['unpaid_salaries']])

        # ...and it keeps showing next month, on the worker's own account page
        first, last = svc.month_bounds(ym)
        next_ym = svc.ym_of(last + timedelta(days=1))
        acc = self.client.get(f'/api/workers/{self.cashier.id}/account/?month={next_ym}').data
        self.assertEqual(acc['owed_by_company'], net)
        self.assertEqual([m['month'] for m in acc['unpaid_months']], [ym])

        # marking it paid clears it everywhere
        line_id = next(l['line_id'] for l in self.client.get(f'/api/salary/{ym}/').data['lines'] if l['worker']['id'] == self.cashier.id)
        self.assertEqual(self.client.post(f'/api/salary-lines/{line_id}/paid/', {'paid': True}, format='json').status_code, 200)
        acc2 = self.client.get(f'/api/workers/{self.cashier.id}/account/?month={next_ym}').data
        self.assertEqual(acc2['owed_by_company'], 0)
        w2 = self.client.get(f'/api/workers/{self.cashier.id}/').data
        self.assertEqual(w2['company_owes'], 0)

    def test_manual_shortage_needs_reason(self):
        r = self.client.post(f'/api/workers/{self.cook.id}/shortages/', {'amount': 500}, format='json')
        self.assertEqual(r.status_code, 400)
        r = self.client.post(f'/api/workers/{self.cook.id}/shortages/', {'amount': 500, 'reason': 'Broken cups'}, format='json')
        self.assertEqual(r.status_code, 201)

    def test_day_and_month_totals_for_money_page(self):
        ym = svc.ym_of(self.today)
        self.sales()
        self.client.post('/api/expenses/', {'category': 'gesi', 'amount': 25000, 'reason': 'gas', 'paid_from': 'simu', 'date': self.d}, format='json')
        self.client.post('/api/advances/', {'worker_id': self.cook.id, 'amount': 7000, 'date': self.d}, format='json')
        day = self.client.get(f'/api/days/{self.d}/').data
        self.assertEqual([(a['worker_name'], a['amount']) for a in day['advances']], [('Juma', 7000)])

        # an expense on a day with no sales still shows up, so the expense column adds up to the month total
        first, _ = svc.month_bounds(ym)
        other = next(first + timedelta(days=i) for i in range(28) if first + timedelta(days=i) != self.today)
        self.client.post('/api/expenses/', {'category': 'gesi', 'amount': 4000, 'reason': 'gas', 'paid_from': 'simu', 'date': other.isoformat()}, format='json')
        r = self.client.get(f'/api/reports/month/{ym}/').data
        rows = {x['date']: x for x in r['days']}
        self.assertEqual((rows[self.d]['expenses'], rows[self.d]['total']), (25000, 350000))
        self.assertEqual((rows[other.isoformat()]['expenses'], rows[other.isoformat()]['total']), (4000, 0))
        self.assertEqual(sum(x['expenses'] for x in r['days']), r['expenses_total'])

    def test_one_day_pay_and_report(self):
        self.sales()
        self.mark_all()  # cashier present, cook late
        self.client.post('/api/expenses/', {'category': 'gesi', 'amount': 25000, 'reason': 'gas', 'paid_from': 'simu'}, format='json')
        self.client.post('/api/advances/', {'worker_id': self.cook.id, 'amount': 7000}, format='json')
        self.client.post(f'/api/workers/{self.cashier.id}/shortages/', {'amount': 2000, 'reason': 'Broken cups'}, format='json')
        self.client.put(f'/api/days/{self.d}/payments/', [{'worker_id': self.cook.id, 'amount': 9000, 'paid_from': 'simu'}], format='json')

        pay = self.client.get(f'/api/days/{self.d}/pay/').data
        rows = {l['worker']['id']: l for l in pay['lines']}
        self.assertEqual((rows[self.cashier.id]['earned'], rows[self.cashier.id]['shortages']), (10000, 2000))  # 300,000 / 30
        self.assertEqual((rows[self.cook.id]['earned'], rows[self.cook.id]['advances']), (9000, 7000))  # what he was handed, not his rate
        self.assertEqual((pay['total_earned'], pay['total_advances'], pay['total_shortages']), (19000, 7000, 2000))

        # an absent monthly worker earns nothing that day
        self.client.put(f'/api/days/{self.d}/attendance/', {str(self.cashier.id): 'absent'}, format='json')
        rows = {l['worker']['id']: l for l in self.client.get(f'/api/days/{self.d}/pay/').data['lines']}
        self.assertEqual(rows[self.cashier.id]['earned'], 0)

    def test_day_report_adds_up_from_the_first(self):
        # two days at the start of last month, so the figures don't depend on today's date
        first = svc.month_bounds(svc.ym_of(self.today.replace(day=1) - timedelta(days=1)))[0]
        d1, d2 = first.isoformat(), (first + timedelta(days=1)).isoformat()
        for d, banda, spend in ((d1, 100000, 20000), (d2, 60000, 5000)):
            self.client.put(f'/api/days/{d}/sales/', {'banda': {'cash': banda, 'mobile': 0}, 'mgahawa': {'cash': 40000, 'mobile': 10000}}, format='json')
            self.client.post('/api/expenses/', {'category': 'gesi', 'amount': spend, 'reason': 'gas', 'paid_from': 'simu', 'date': d}, format='json')
        salaries = self.client.get(f'/api/reports/month/{svc.ym_of(first)}/').data['salaries_total']

        r = self.client.get(f'/api/reports/day/{d1}/').data
        self.assertEqual((r['total'], r['expenses_total'], r['profit']), (150000, 20000, 150000 - 20000 - salaries))

        # on the 2nd: sales of the 1st + 2nd, expenses of the 1st + 2nd, minus the whole month's salaries
        r = self.client.get(f'/api/reports/day/{d2}/').data
        self.assertEqual((r['from'], r['total'], r['expenses_total'], r['salaries_total']), (d1, 260000, 25000, salaries))
        self.assertEqual(r['profit'], 260000 - 25000 - salaries)
        self.assertEqual((r['day_sales'], r['day_expenses']), (110000, 5000))
        self.assertEqual([(x['date'], x['sales_to_date'], x['expenses_to_date']) for x in r['running']], [(d1, 150000, 20000), (d2, 260000, 25000)])

    def test_daily_payments(self):
        ym = svc.ym_of(self.today)
        self.sales()  # banda: 100,000 cash
        day = self.client.get(f'/api/days/{self.d}/').data
        self.assertEqual([x['worker_id'] for x in day['payments']], [self.cook.id])  # only daily-paid workers are listed

        # a payment from the banda till comes out of the banda's expected cash
        r = self.client.put(f'/api/days/{self.d}/payments/', [{'worker_id': self.cook.id, 'amount': 8000, 'paid_from': 'droo', 'section': 'banda'}], format='json')
        self.assertEqual(r.status_code, 200, r.data)
        b = r.data['sections']['banda']
        self.assertEqual((r.data['payments_total'], b['wages'], b['expected']), (8000, 8000, 100000 - 8000))
        bad = self.client.put(f'/api/days/{self.d}/payments/', [{'worker_id': self.cook.id, 'amount': 8000, 'paid_from': 'droo'}], format='json')
        self.assertEqual(bad.status_code, 400)  # a till payment must say which till

        # not on the monthly salary list; shown with the daily-paid workers instead
        sal = self.client.get(f'/api/salary/{ym}/').data
        self.assertNotIn(self.cook.id, [l['worker']['id'] for l in sal['lines']])
        self.assertEqual([(x['worker']['id'], x['days'], x['total']) for x in sal['daily']], [(self.cook.id, 1, 8000)])
        acct = self.client.get(f'/api/workers/{self.cook.id}/account/').data
        self.assertEqual((acct['pay_type'], acct['days_paid'], acct['base'], acct['net']), ('daily', 1, 8000, 0))

        # the reports subtract it on the day it was paid
        rep = self.client.get(f'/api/reports/day/{self.d}/').data
        self.assertEqual((rep['payments_total'], rep['day_payments']), (8000, 8000))
        self.assertEqual(rep['profit'], rep['total'] - rep['expenses_total'] - 8000 - rep['salaries_total'])
        month = self.client.get(f'/api/reports/month/{ym}/').data
        self.assertEqual(month['payments_total'], 8000)
        self.assertEqual(month['profit'], month['total'] - month['expenses_total'] - 8000 - month['salaries_total'])

        # an empty amount removes it; once the day is closed it can't change
        r = self.client.put(f'/api/days/{self.d}/payments/', [{'worker_id': self.cook.id, 'amount': None}], format='json')
        self.assertEqual(r.data['payments_total'], 0)
        self.client.put(f'/api/days/{self.d}/cash-count/', {'banda': 50000 + 100000, 'mgahawa': 80000 + 200000}, format='json')
        self.mark_all()
        self.assertEqual(self.client.post(f'/api/days/{self.d}/close/').status_code, 200)
        r = self.client.put(f'/api/days/{self.d}/payments/', [{'worker_id': self.cook.id, 'amount': 5000, 'paid_from': 'simu'}], format='json')
        self.assertEqual((r.status_code, r.data['code']), (409, 'day_closed'))

    def test_switch_pay_type(self):
        ym = svc.ym_of(self.today)
        r = self.client.patch(f'/api/workers/{self.cashier.id}/', {'pay_type': 'daily', 'rate': 12000}, format='json')
        self.assertEqual((r.status_code, r.data['pay_type'], r.data['rate']), (200, 'daily', 12000))
        sal = self.client.get(f'/api/salary/{ym}/').data
        self.assertNotIn(self.cashier.id, [l['worker']['id'] for l in sal['lines']])
        self.assertIn(self.cashier.id, [x['worker']['id'] for x in sal['daily']])
        self.assertIn(self.cashier.id, [x['worker_id'] for x in self.client.get(f'/api/days/{self.d}/').data['payments']])
        log = AuditLog.objects.get(action='worker.pay')
        self.assertEqual((log.detail['pay_type'], log.detail['was']), ('daily', {'pay_type': 'monthly', 'rate': 300000}))
        self.assertEqual(self.client.patch(f'/api/workers/{self.cashier.id}/', {'pay_type': 'weekly'}, format='json').status_code, 400)
        self.assertEqual(self.client.patch(f'/api/workers/{self.cashier.id}/', {'rate': 0}, format='json').status_code, 400)

    def test_reports_show_daily_pay(self):
        from io import BytesIO
        from unittest import mock

        from openpyxl import load_workbook
        from reportlab import rl_config
        ym = svc.ym_of(self.today)
        self.sales()
        self.client.put(f'/api/days/{self.d}/payments/', [{'worker_id': self.cook.id, 'amount': 8000, 'paid_from': 'droo', 'section': 'banda'}], format='json')
        rep = self.client.get(f'/api/reports/day/{self.d}/').data
        money_left = f'{rep["profit"]:,}' if rep['profit'] >= 0 else f'-{-rep["profit"]:,}'
        with mock.patch.object(rl_config, 'pageCompression', 0):  # readable PDF text
            day_pdf = self.client.get(f'/api/reports/day/{self.d}/?file=pdf&lang=en').content
            month_pdf = self.client.get(f'/api/reports/month/{ym}/?file=pdf&lang=en').content
        for pdf in (day_pdf, month_pdf):
            self.assertIn(b'Daily pay', pdf)
            self.assertIn(b'8,000', pdf)
            self.assertIn(rb'Money left \(profit\)', pdf)  # parentheses are escaped inside PDF strings
        self.assertIn(money_left.encode(), day_pdf)
        self.assertIn(b'Juma', day_pdf)  # who was paid

        book = load_workbook(BytesIO(self.client.get(f'/api/reports/day/{self.d}/?file=xlsx&lang=en').content))
        cells = [c for ws in book.worksheets for row in ws.iter_rows(values_only=True) for c in row]
        self.assertIn('Total daily pay', cells)
        self.assertIn(8000, cells)
        self.assertIn(rep['profit'], cells)
        book = load_workbook(BytesIO(self.client.get(f'/api/reports/month/{ym}/?file=xlsx&lang=en').content))
        self.assertIn('Daily pay', book.sheetnames)
        self.assertEqual([r[1:] for r in book['Daily pay'].iter_rows(min_row=2, max_row=2, values_only=True)], [('Juma', 'Till - Stall', 8000)])

    def test_salary_formulas(self):
        ym = svc.ym_of(self.today)
        first, _ = svc.month_bounds(ym)
        # fixed history on the month's first working days: the cook late twice then present, the cashier absent
        # twice then present. Their shared day off is skipped, since a day off overrides any mark and is neither
        # paid (daily) nor cut (monthly) under the default rules; otherwise this test fails on some dates.
        marked = []
        for i in range(10):
            d = first + timedelta(days=i)
            if d > self.today:
                break
            if svc.js_dow(d) == self.cook.day_off:
                continue
            n = len(marked)
            Attendance.objects.update_or_create(worker=self.cook, date=d, defaults={'status': 'late' if n < 2 else 'present'})
            Attendance.objects.update_or_create(worker=self.cashier, date=d, defaults={'status': 'absent' if n < 2 else 'present'})
            DailyPayment.objects.create(worker=self.cook, date=d, amount=8000 if n == 0 else 10000, paid_from='simu')
            marked.append(d)
        f = svc.salary_figures(self.cook, ym, Settings.load().rules)
        self.assertEqual((f['base'], f['net']), (10000 * len(marked) - (2000 if marked else 0), 0))
        g = svc.salary_figures(self.cashier, ym, Settings.load().rules)
        absent = min(2, len(marked))
        self.assertEqual(g['base'], 300000 - round(300000 / 30 * absent))


class OpinionsTest(APITestCase):
    def setUp(self):
        cache.clear()  # the spam guard counts in the cache
        self.manager = get_user_model().objects.create_user('m', password='x')

    def send(self, **body):
        return self.client.post('/api/opinions/submit/', {'source': 'customer', 'message': 'Chakula kitamu sana', **body}, format='json')

    def test_anyone_can_send_and_it_stays_anonymous(self):
        r = self.send(rating=5, topic='food')
        self.assertEqual(r.status_code, 201, r.data)
        o = Opinion.objects.get()
        self.assertEqual((o.source, o.topic, o.rating, o.contact, o.read_at), ('customer', 'food', 5, '', None))
        # a worker may leave a name if they want a reply; topic and rating can be left out
        r = self.send(source='worker', message='  Tunahitaji gesi zaidi  ', contact='Juma 0712 000 000', rating='', topic=None)
        self.assertEqual(r.status_code, 201, r.data)
        o = Opinion.objects.get(source='worker')
        self.assertEqual((o.topic, o.rating, o.message, o.contact), ('other', None, 'Tunahitaji gesi zaidi', 'Juma 0712 000 000'))
        # nothing about the sender's device is saved
        self.assertEqual({f.name for f in Opinion._meta.fields},
                         {'id', 'created_at', 'source', 'topic', 'rating', 'message', 'contact', 'read_at'})

    def test_form_rules(self):
        for body, field in [({'source': 'boss'}, 'source'), ({'message': 'ok'}, 'message'), ({'message': 'x' * 1001}, 'message'),
                            ({'rating': 6}, 'rating'), ({'rating': 0}, 'rating'), ({'topic': 'music'}, 'topic'),
                            ({'contact': 'x' * 101}, 'contact')]:
            r = self.send(**body)
            self.assertEqual(r.status_code, 400, body)
            self.assertIn(field, r.data['errors'])
        self.assertEqual(self.send(message='x' * 1000).status_code, 201)

    def test_bots_filling_the_hidden_field_are_ignored(self):
        self.assertEqual(self.send(website='http://spam.example').status_code, 201)
        self.assertEqual(Opinion.objects.count(), 0)

    def test_one_device_cannot_flood_the_form(self):
        for _ in range(20):
            self.assertEqual(self.send().status_code, 201)
        self.assertEqual(self.send().status_code, 429)
        self.assertEqual(Opinion.objects.count(), 20)

    def test_an_expired_login_does_not_block_the_form(self):
        self.client.credentials(HTTP_AUTHORIZATION='Bearer not-a-valid-token')
        self.assertEqual(self.send().status_code, 201)

    def test_only_the_manager_reads_them(self):
        self.send(rating=4)
        self.send(source='worker', rating=2)
        self.send(source='worker')
        for path in ('/api/opinions/', '/api/opinions/unread/'):
            self.assertEqual(self.client.get(path).status_code, 401)
        self.client.force_authenticate(self.manager)
        r = self.client.get('/api/opinions/')
        self.assertEqual(len(r.data['opinions']), 3)
        self.assertEqual(r.data['counts'], {'total': 3, 'unread': 3, 'customer': 1, 'worker': 2, 'rated': 2, 'average_rating': 3.0})
        first = r.data['opinions'][0]['id']
        r = self.client.patch(f'/api/opinions/{first}/', {'read': True}, format='json')
        self.assertTrue(r.data['read'])
        self.assertEqual(self.client.get('/api/opinions/unread/').data, {'unread': 2})
        self.assertEqual(self.client.patch(f'/api/opinions/{first}/', {'read': 'yes'}, format='json').status_code, 400)
        self.client.patch(f'/api/opinions/{first}/', {'read': False}, format='json')
        self.assertEqual(self.client.post('/api/opinions/read-all/').data, {'marked': 3})
        self.assertEqual(self.client.get('/api/opinions/unread/').data, {'unread': 0})
        self.assertEqual(self.client.delete(f'/api/opinions/{first}/').status_code, 204)
        self.assertEqual(Opinion.objects.count(), 2)

    def test_inbox_pages_and_filters(self):
        for i in range(5):
            self.send(source='worker' if i % 2 else 'customer', message=f'Maoni namba {i}')
        self.client.force_authenticate(self.manager)
        r = self.client.get('/api/opinions/?limit=2')
        self.assertEqual([o['message'] for o in r.data['opinions']], ['Maoni namba 4', 'Maoni namba 3'])  # newest first
        self.assertTrue(r.data['has_more'])
        self.assertEqual(r.data['counts']['total'], 5)  # totals always cover everything
        r = self.client.get('/api/opinions/?source=worker&limit=10')
        self.assertEqual(len(r.data['opinions']), 2)
        self.assertFalse(r.data['has_more'])
        self.client.patch(f"/api/opinions/{r.data['opinions'][0]['id']}/", {'read': True}, format='json')
        self.assertEqual(len(self.client.get('/api/opinions/?unread=1').data['opinions']), 4)
        self.assertEqual(len(self.client.get('/api/opinions/').data['opinions']), 5)  # no limit: everything, as before


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])  # fast hashing: many logins here
class SecurityTest(APITestCase):
    PW = 'Siri-Ndefu-2026!'

    def setUp(self):
        cache.clear()  # login and spam limits count in the cache
        self.user = get_user_model().objects.create_user('meneja', password=self.PW)
        ExpenseCategory.objects.create(key='gesi', name_sw='Gesi', name_en='Gas')

    def login(self, password=None, **extra):
        return self.client.post('/api/auth/login/', {'username': 'meneja', 'password': password or self.PW}, format='json', **extra)

    def bearer(self, access):
        c = self.client_class()
        c.credentials(HTTP_AUTHORIZATION=f'Bearer {access}')
        return c

    def test_password_guessing_is_slowed_down(self):
        codes = [self.login(f'guess{i}').status_code for i in range(11)]
        self.assertEqual(codes[:10], [401] * 10)
        self.assertEqual(codes[10], 429)  # the 11th try in a minute from one device
        cache.clear()
        # many devices guessing at one account: blocked after 30 an hour
        codes = [self.login(f'guess{i}', REMOTE_ADDR=f'10.0.{i // 250}.{i % 250 + 1}').status_code for i in range(31)]
        self.assertEqual((codes.count(401), codes[-1]), (30, 429))

    def test_new_password_ends_every_other_sign_in(self):
        old = self.login().data
        r = self.bearer(old['access']).post('/api/auth/change-password/', {'current_password': self.PW, 'new_password': 'Mpya-Kabisa-2026?'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(self.bearer(old['access']).get('/api/auth/me/').status_code, 401)  # a stolen token stops working
        self.assertEqual(self.client.post('/api/auth/refresh/', {'refresh': old['refresh']}, format='json').status_code, 401)
        self.assertEqual(self.bearer(r.data['access']).get('/api/auth/me/').status_code, 200)  # this device stays in
        self.assertEqual(self.client.post('/api/auth/refresh/', {'refresh': r.data['refresh']}, format='json').status_code, 200)

    def test_refresh_tokens_work_once_and_sign_out_ends_them(self):
        t = self.login().data
        self.assertEqual(self.client.post('/api/auth/refresh/', {'refresh': t['refresh']}, format='json').status_code, 200)
        self.assertEqual(self.client.post('/api/auth/refresh/', {'refresh': t['refresh']}, format='json').status_code, 401)
        t = self.login().data
        self.assertEqual(self.client.post('/api/auth/logout/', {'refresh': t['refresh']}, format='json').status_code, 204)
        self.assertEqual(self.client.post('/api/auth/refresh/', {'refresh': t['refresh']}, format='json').status_code, 401)
        self.assertEqual(self.client.post('/api/auth/logout/', {'refresh': 'rubbish'}, format='json').status_code, 204)

    def test_a_made_up_address_cannot_dodge_the_spam_limit(self):
        with override_settings(REST_FRAMEWORK={**django_settings.REST_FRAMEWORK, 'NUM_PROXIES': 1}):
            send = lambda i: self.client.post(  # noqa: E731  the proxy appends the real address after anything faked
                '/api/opinions/submit/', {'source': 'customer', 'message': 'spam spam'}, format='json',
                HTTP_X_FORWARDED_FOR=f'10.9.{i}.1, 203.0.113.7')
            codes = [send(i).status_code for i in range(21)]
        self.assertEqual((codes.count(201), codes[-1]), (20, 429))

    def test_the_whole_form_has_a_daily_ceiling(self):
        with mock.patch('core.views.OpinionDailyCap.rate', '3/day'):
            codes = [self.client.post('/api/opinions/submit/', {'source': 'worker', 'message': 'maoni mengi'}, format='json',
                                      REMOTE_ADDR=f'10.1.1.{i + 1}').status_code for i in range(4)]
        self.assertEqual(codes, [201, 201, 201, 429])

    def test_absurd_amounts_are_refused_not_crashed(self):
        c = self.bearer(self.login().data['access'])
        r = c.post('/api/expenses/', {'category': 'gesi', 'amount': 10 ** 30, 'reason': 'x', 'paid_from': 'simu'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('amount', r.data['errors'])

    def test_excel_never_runs_text_as_a_formula(self):
        from openpyxl import Workbook
        from core import exports
        ws = Workbook().active
        exports._sheet(ws, ['Jina'], [['=HYPERLINK("http://evil.example","bonyeza")'], ['Amina']], [20])
        self.assertEqual([c.data_type for c in ws['A']], ['s', 's', 's'])
        self.assertEqual(ws['A2'].value, '=HYPERLINK("http://evil.example","bonyeza")')  # shown exactly as typed


class ContentSecurityPolicyTest(APITestCase):
    def test_the_inline_script_is_allowed_by_the_hash_browsers_compute(self):
        import base64
        import hashlib
        from config import csp
        with tempfile.TemporaryDirectory() as d:
            Path(d, 'index.html').write_bytes(b'<html><script>\r\n  var t = 1;\r\n</script></html>')  # Windows line ends
            with override_settings(FRONTEND_DIST=d):
                csp.policy.cache_clear()
                policy = csp.policy()
            csp.policy.cache_clear()
        browser_hash = base64.b64encode(hashlib.sha256(b'\n  var t = 1;\n').digest()).decode()  # browsers see LF only
        self.assertIn(f"script-src 'self' 'sha256-{browser_hash}'", policy)
        self.assertIn("frame-ancestors 'none'", policy)


class BackupTest(TransactionTestCase):
    # Not wrapped in a test transaction: SQLite's backup waits for open write transactions to finish.
    def test_backup_copies_the_data_and_keeps_the_newest(self):
        Worker.objects.create(name='Rehema', role='keshia', pay_type='monthly', rate=300000, joined_on=svc.today())
        with tempfile.TemporaryDirectory() as tmp:
            for name in ('daftari-2026-01-01_000000.sqlite3', 'daftari-2026-01-02_000000.sqlite3'):
                Path(tmp, name).write_bytes(b'')
            call_command('backup_db', dir=tmp, keep=2, stdout=io.StringIO())
            files = sorted(p.name for p in Path(tmp).glob('daftari-*.sqlite3'))
            self.assertEqual(len(files), 2)
            self.assertNotIn('daftari-2026-01-01_000000.sqlite3', files)  # the oldest was removed
            with closing(sqlite3.connect(Path(tmp, files[-1]))) as db:
                self.assertEqual(db.execute('select name from core_worker').fetchall(), [('Rehema',)])
