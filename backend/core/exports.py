"""Branded PDF (reportlab) and Excel (openpyxl) reports."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle, Paragraph
from reportlab.lib.styles import ParagraphStyle

EMERALD, BRASS, INK, MUTED, PAPER = '#0b5d49', '#c8962e', '#16231f', '#6b7a74', '#f4f7f4'
DAYS = {'sw': ['Jumapili', 'Jumatatu', 'Jumanne', 'Jumatano', 'Alhamisi', 'Ijumaa', 'Jumamosi'],
        'en': ['Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']}
DAYS_SHORT = {'sw': ['Jpl', 'Jtt', 'Jnn', 'Jtn', 'Alh', 'Ijm', 'Jms'], 'en': ['Sun', 'Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat']}
MONTHS = {'sw': ['Januari', 'Februari', 'Machi', 'Aprili', 'Mei', 'Juni', 'Julai', 'Agosti', 'Septemba', 'Oktoba', 'Novemba', 'Desemba'],
          'en': ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']}


def n(v):
    return f'{round(v or 0):,}'


def signed(v):
    return '0' if not v else f'{"+" if v > 0 else "-"}{n(abs(v))}'


def T(lang):
    return lambda sw, en: sw if lang == 'sw' else en


def month_title(ym, lang):
    y, m = ym.split('-')
    return f'{MONTHS[lang][int(m) - 1]} {y}'


def short(iso, lang):
    d = date.fromisoformat(iso)
    return f'{d.day} {MONTHS[lang][d.month - 1][:3]}'


# ------------------------------------------------------------------ PDF
def _doc(buf, title):
    return SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm,
                             bottomMargin=16 * mm, title=title, author='Daftari')


def _header(title, subtitle):
    h1 = ParagraphStyle('h1', fontName='Helvetica-Bold', fontSize=20, textColor=colors.white, leading=24)
    sub = ParagraphStyle('sub', fontName='Helvetica', fontSize=10, textColor=colors.HexColor('#cfe5dc'), leading=14)
    brand = ParagraphStyle('brand', fontName='Helvetica-Bold', fontSize=9, textColor=colors.HexColor(BRASS), leading=12)
    t = Table([[Paragraph('DAFTARI', brand)], [Paragraph(title, h1)], [Paragraph(subtitle, sub)]], colWidths=[178 * mm])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(EMERALD)), ('LEFTPADDING', (0, 0), (-1, -1), 14),
        ('TOPPADDING', (0, 0), (0, 0), 12), ('BOTTOMPADDING', (0, -1), (-1, -1), 14),
        ('LINEBELOW', (0, -1), (-1, -1), 3, colors.HexColor(BRASS)),
    ]))
    return t


def _kpis(items):
    lab = ParagraphStyle('kl', fontName='Helvetica', fontSize=8, textColor=colors.HexColor(MUTED), leading=10)
    val = ParagraphStyle('kv', fontName='Helvetica-Bold', fontSize=12, textColor=colors.HexColor(INK), leading=15)
    cells = [[Paragraph(label.upper(), lab), Paragraph(value, val)] for label, value in items]
    row = [Table([[c[0]], [c[1]]], colWidths=[(178 * mm) / len(items) - 4]) for c in cells]
    t = Table([row], colWidths=[(178 * mm) / len(items)] * len(items))
    t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(PAPER)),
                           ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 8)]))
    return t


def _table(rows, widths, right_from=1, total_last=False):
    t = Table(rows, colWidths=widths, repeatRows=1)
    style = [
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'), ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(EMERALD)), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('ALIGN', (right_from, 0), (-1, -1), 'RIGHT'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor(PAPER)]),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LINEBELOW', (0, 0), (-1, -1), 0.25, colors.HexColor('#dde5e0')),
    ]
    if total_last:
        style += [('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'), ('LINEABOVE', (0, -1), (-1, -1), 1.2, colors.HexColor(BRASS))]
    t.setStyle(TableStyle(style))
    return t


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('Helvetica', 7.5)
    canvas.setFillColor(colors.HexColor(MUTED))
    canvas.drawString(16 * mm, 9 * mm, 'Daftari')
    canvas.drawRightString(194 * mm, 9 * mm, str(doc.page))
    canvas.restoreState()


def _build(title, story):
    buf = BytesIO()
    _doc(buf, title).build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


def _h2(text):
    return Paragraph(text, ParagraphStyle('h2', fontName='Helvetica-Bold', fontSize=11.5, textColor=colors.HexColor(EMERALD),
                                          leading=15, spaceAfter=4))


def _paid_from(x, lang):
    t = T(lang)
    label = {'droo': t('Droo', 'Till'), 'simu': t('Pesa za simu', 'Mobile money'), 'other': t('Nyingine', 'Other')}[x['paid_from']]
    if x.get('section'):
        label += ' - ' + {'banda': t('Banda', 'Stall'), 'mgahawa': t('Mgahawa', 'Restaurant')}[x['section']]
    return label


def _profit_rows(lang, sales, spent, paid, salaries, profit, span=''):
    """The profit sum, line by line, exactly as the Ripoti screen works it out."""
    t = T(lang)
    span = f' ({span})' if span else ''
    return [[t('Hesabu ya faida', 'Profit'), 'TSh'],
            [t('Mauzo', 'Sales') + span, n(sales)],
            [t('- Matumizi', '- Expenses') + span, n(spent)],
            [t('- Malipo ya siku', '- Daily pay') + span, n(paid)],
            [t('- Mishahara ya mwezi mzima', "- The whole month's salaries"), n(salaries)],
            [t('Pesa iliyobaki (faida)', 'Money left (profit)'), signed(profit) if profit < 0 else n(profit)]]


def month_pdf(r, lang):
    t = T(lang)
    title = t('Ripoti ya mauzo ya mwezi', 'Monthly sales report')
    best = r['best']
    story = [_header(title, month_title(r['month'], lang)), Spacer(1, 6 * mm), _kpis([
        (t('Jumla ya mauzo', 'Total sales'), f'TSh {n(r["total"])}'), (t('Pesa taslimu', 'Cash'), f'TSh {n(r["cash"])}'),
        (t('Pesa za simu', 'Mobile money'), f'TSh {n(r["mobile"])}'), (t('Wastani/siku', 'Avg / day'), f'TSh {n(r["average"])}')]),
        Spacer(1, 3 * mm), _kpis([
            (t('Banda', 'Stall'), f'TSh {n(r["banda"])}'), (t('Mgahawa', 'Restaurant'), f'TSh {n(r["mgahawa"])}'),
            (t('Siku bora', 'Best day'), f'{short(best["date"], lang)}: {n(best["total"])}' if best else '-'),
            (t('Pesa iliyobaki (faida)', 'Money left (profit)'), f'TSh {signed(r["profit"]) if r["profit"] < 0 else n(r["profit"])}')]), Spacer(1, 6 * mm),
        _table(_profit_rows(lang, r['total'], r['expenses_total'], r['payments_total'], r['salaries_total'], r['profit']),
               [120 * mm, 58 * mm], 1, True), Spacer(1, 6 * mm),
        _h2(t('Kila siku', 'Every day'))]
    rows = [[t('Tarehe', 'Date'), t('Siku', 'Day'), t('Banda', 'Stall'), t('Mgahawa', 'Restaurant'), t('Mauzo', 'Sales'),
             t('Matumizi', 'Expenses'), t('Malipo', 'Daily pay'), t('Tofauti', 'Diff.')]]
    for d in sorted(r['days'], key=lambda x: x['date']):
        rows.append([short(d['date'], lang), DAYS_SHORT[lang][d['dow']], n(d['banda']['cash'] + d['banda']['mobile']),
                     n(d['mgahawa']['cash'] + d['mgahawa']['mobile']), n(d['total']), n(d['expenses']), n(d['payments']), signed(d['diff'])])
    rows.append([t('JUMLA', 'TOTAL'), '', n(r['banda']), n(r['mgahawa']), n(r['total']), n(r['expenses_total']),
                 n(r['payments_total']), signed(r['diff_total'])])
    story.append(_table(rows, [20 * mm, 14 * mm, 24 * mm, 24 * mm, 26 * mm, 24 * mm, 23 * mm, 23 * mm], 2, True))
    paid = [x for x in r['daily_paid'] if x['days']]
    if paid:
        rows = [[t('Malipo ya kila siku', 'Daily pay'), t('Siku', 'Days'), 'TSh']]
        rows += [[x['worker']['name'], str(x['days']), n(x['total'])] for x in paid]
        rows.append([t('JUMLA', 'TOTAL'), str(sum(x['days'] for x in paid)), n(r['payments_total'])])
        story += [Spacer(1, 6 * mm), _table(rows, [110 * mm, 30 * mm, 38 * mm], 1, True)]
    if r['expenses_by_category']:
        story += [Spacer(1, 6 * mm), _table(
            [[t('Matumizi kwa kundi', 'Expenses by category'), 'TSh']] + [
                [c['name_sw' if lang == 'sw' else 'name_en'], n(c['amount'])] for c in r['expenses_by_category']],
            [120 * mm, 58 * mm], 1)]
    return _build(title, story)


def day_pdf(p, rep, lang):
    t = T(lang)
    d = date.fromisoformat(p['date'])
    title = t('Ripoti ya siku', 'Daily report')
    sub = f'{DAYS[lang][p["dow"]]}, {d.day} {MONTHS[lang][d.month - 1]} {d.year}'
    S = p['sections']
    left = p['sales_total'] - p['expenses_total'] - p['payments_total']  # what the day itself left, before monthly salaries
    rows = [[t('Sehemu', 'Section'), t('Pesa taslimu', 'Cash'), t('Pesa za simu', 'Mobile money'), t('Jumla', 'Total')]]
    for key, name in (('banda', t('Banda', 'Stall')), ('mgahawa', t('Mgahawa', 'Restaurant'))):
        rows.append([name, n(S[key]['cash']), n(S[key]['mobile']), n(S[key]['cash'] + S[key]['mobile'])])
    rows.append([t('JUMLA', 'TOTAL'), n(p['cash_total']), n(p['mobile_total']), n(p['sales_total'])])
    story = [_header(title, sub), Spacer(1, 6 * mm), _kpis([
        (t('Mauzo', 'Sales'), f'TSh {n(p["sales_total"])}'), (t('Matumizi', 'Expenses'), f'TSh {n(p["expenses_total"])}'),
        (t('Malipo ya siku', 'Daily pay'), f'TSh {n(p["payments_total"])}'),
        (t('Baki ya siku', 'Left this day'), f'TSh {signed(left) if left < 0 else n(left)}'),
        (t('Tofauti ya pesa', 'Cash difference'), f'TSh {signed(p["difference_total"])}')]), Spacer(1, 6 * mm),
        _table(rows, [58 * mm, 40 * mm, 40 * mm, 40 * mm], 1, True), Spacer(1, 6 * mm)]
    exp = [[t('Matumizi', 'Expense'), t('Maelezo', 'Reason'), 'TSh']] + [
        [e['category_sw' if lang == 'sw' else 'category_en'], e['reason'], n(e['amount'])] for e in p['expenses']]
    if len(exp) == 1:
        exp.append([t('Hakuna matumizi', 'No expenses'), '', '0'])
    story += [_table(exp, [45 * mm, 100 * mm, 33 * mm], 2), Spacer(1, 6 * mm)]
    paid = [x for x in p['payments'] if x['amount']]
    pay = [[t('Malipo ya siku', 'Daily pay'), t('Imetoka', 'Paid from'), 'TSh']]
    pay += [[x['name'], _paid_from(x, lang), n(x['amount'])] for x in paid]
    pay.append([t('JUMLA', 'TOTAL'), '', n(p['payments_total'])] if paid else [t('Hakuna malipo ya siku', 'No daily pay'), '', '0'])
    story += [_table(pay, [75 * mm, 70 * mm, 33 * mm], 2, bool(paid)), Spacer(1, 8 * mm)]

    # the Ripoti figures: sales, expenses and daily pay added up from the 1st, minus the month's salaries
    span = f'{short(rep["from"], lang)} - {short(rep["date"], lang)}'
    story += [_h2(t(f'Pesa iliyobaki (faida) hadi siku hii ({span})', f'Money left (profit) up to this day ({span})')),
              _table(_profit_rows(lang, rep['total'], rep['expenses_total'], rep['payments_total'], rep['salaries_total'], rep['profit'], span),
                     [120 * mm, 58 * mm], 1, True)]
    if rep['running']:
        run = [[t('Tarehe', 'Date'), t('Mauzo', 'Sales'), t('Matumizi', 'Expenses'), t('Malipo', 'Daily pay'),
                t('Mauzo jumla', 'Sales so far'), t('Matumizi jumla', 'Expenses so far'), t('Malipo jumla', 'Pay so far')]]
        run += [[short(x['date'], lang), n(x['sales']), n(x['expenses']), n(x['payments']),
                 n(x['sales_to_date']), n(x['expenses_to_date']), n(x['payments_to_date'])] for x in rep['running']]
        story += [Spacer(1, 6 * mm), _h2(t('Jinsi yalivyoongezeka', 'How it added up')),
                  _table(run, [20 * mm, 25 * mm, 25 * mm, 24 * mm, 29 * mm, 29 * mm, 26 * mm], 1)]
    return _build(title, story)


# ------------------------------------------------------------------ Excel
def _sheet(ws, header, rows, widths, total=False):
    ws.append(header)
    for row in rows:
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=EMERALD.lstrip('#'))
        c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = '#,##0;[Red]-#,##0'
    if total:
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
            c.fill = PatternFill('solid', fgColor='F3E6C4')
    ws.freeze_panes = 'A2'


def _xlsx(wb):
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _profit_sheet(wb, lang, sales, spent, paid, salaries, profit, span=''):
    t = T(lang)
    span = f' ({span})' if span else ''
    ws = wb.create_sheet(t('Faida', 'Profit'))
    _sheet(ws, [t('Hesabu ya faida', 'Profit'), 'TSh'], [
        [t('Mauzo', 'Sales') + span, sales], [t('Matumizi', 'Expenses') + span, -spent],
        [t('Malipo ya siku', 'Daily pay') + span, -paid], [t('Mishahara ya mwezi mzima', "The whole month's salaries"), -salaries],
        [t('Pesa iliyobaki (faida)', 'Money left (profit)'), profit]], [44, 18], total=True)


def month_xlsx(r, lang):
    t = T(lang)
    wb = Workbook()
    ws = wb.active
    ws.title = month_title(r['month'], lang)[:31]
    rows = [[d['date'], DAYS[lang][d['dow']], d['banda']['cash'], d['banda']['mobile'], d['mgahawa']['cash'],
             d['mgahawa']['mobile'], d['total'], d['expenses'], d['payments'], d['diff']] for d in sorted(r['days'], key=lambda x: x['date'])]
    rows.append([t('JUMLA', 'TOTAL'), '', *[sum(x[i] for x in rows) for i in range(2, 10)]])
    _sheet(ws, [t('Tarehe', 'Date'), t('Siku', 'Day'), t('Banda: pesa taslimu', 'Stall: cash'),
                t('Banda: pesa za simu', 'Stall: mobile'), t('Mgahawa: pesa taslimu', 'Restaurant: cash'),
                t('Mgahawa: pesa za simu', 'Restaurant: mobile'), t('Jumla ya mauzo', 'Total sales'),
                t('Matumizi', 'Expenses'), t('Malipo ya siku', 'Daily pay'),
                t('Tofauti ya pesa', 'Cash difference')], rows, [14, 14, 18, 18, 20, 20, 18, 16, 16, 16], total=True)
    _profit_sheet(wb, lang, r['total'], r['expenses_total'], r['payments_total'], r['salaries_total'], r['profit'])
    pay = wb.create_sheet(t('Malipo ya siku', 'Daily pay'))
    prow = [[x['date'], x['worker_name'], _paid_from(x, lang), x['amount']] for x in sorted(r['payment_rows'], key=lambda x: (x['date'], x['worker_name']))]
    prow.append([t('JUMLA', 'TOTAL'), '', '', r['payments_total']])
    _sheet(pay, [t('Tarehe', 'Date'), t('Mfanyakazi', 'Worker'), t('Imetoka', 'Paid from'), 'TSh'], prow, [14, 28, 26, 16], total=True)
    return _xlsx(wb)


def day_xlsx(p, rep, lang):
    t = T(lang)
    wb = Workbook()
    ws = wb.active
    ws.title = p['date']
    S = p['sections']
    rows = []
    for key, name in (('banda', t('Banda', 'Stall')), ('mgahawa', t('Mgahawa', 'Restaurant'))):
        rows += [[name, t('Pesa taslimu', 'Cash'), S[key]['cash']], [name, t('Pesa za simu', 'Mobile money'), S[key]['mobile']]]
    rows.append([t('Jumla ya mauzo', 'Total sales'), '', p['sales_total']])
    rows += [[t('Matumizi', 'Expense'), f'{e["category_sw" if lang == "sw" else "category_en"]}: {e["reason"]}', e['amount']]
             for e in p['expenses']]
    rows.append([t('Jumla ya matumizi', 'Total expenses'), '', p['expenses_total']])
    rows += [[t('Malipo ya siku', 'Daily pay'), f'{x["name"]} ({_paid_from(x, lang)})', x['amount']] for x in p['payments'] if x['amount']]
    rows += [[t('Jumla ya malipo ya siku', 'Total daily pay'), '', p['payments_total']],
             [t('Tofauti ya pesa', 'Cash difference'), '', p['difference_total']],
             [t('Baki ya siku (mauzo - matumizi - malipo)', 'Left this day (sales - expenses - daily pay)'), '', p['sales_total'] - p['expenses_total'] - p['payments_total']]]
    _sheet(ws, [t('Sehemu', 'Section'), t('Maelezo', 'Detail'), 'TSh'], rows, [26, 44, 16], total=True)
    span = f'{short(rep["from"], lang)} - {short(rep["date"], lang)}'
    _profit_sheet(wb, lang, rep['total'], rep['expenses_total'], rep['payments_total'], rep['salaries_total'], rep['profit'], span)
    run = wb.create_sheet(t('Kila siku', 'Every day'))
    _sheet(run, [t('Tarehe', 'Date'), t('Mauzo', 'Sales'), t('Matumizi', 'Expenses'), t('Malipo ya siku', 'Daily pay'),
                 t('Mauzo jumla', 'Sales so far'), t('Matumizi jumla', 'Expenses so far'), t('Malipo jumla', 'Pay so far')],
           [[x['date'], x['sales'], x['expenses'], x['payments'], x['sales_to_date'], x['expenses_to_date'], x['payments_to_date']]
            for x in rep['running']], [14, 16, 16, 16, 18, 18, 16])
    return _xlsx(wb)
